"""RAG 摄取管线单测：extract / clean / chunker / store / service。"""
from __future__ import annotations

import asyncio
import io

from myopencode.config.schema import RagConfig
from myopencode.rag.bm25 import BM25Index, tokenize
from myopencode.rag.chunker import Chunk, chunk_text, dedup_chunks, kind_for_ext
from myopencode.rag.clean import clean_text
from myopencode.rag.extract import extract_text, file_ext
from myopencode.rag.reranker import _parse_order
from myopencode.rag.service import RagService
from myopencode.rag.store import RagStore


# ── extract ─────────────────────────────────────────────────────────
def test_file_ext():
    assert file_ext("a.PDF") == ".pdf"
    assert file_ext("README.md") == ".md"
    assert file_ext("noext") == ""


def test_extract_text_and_unknown():
    assert extract_text("a.txt", "你好".encode(), {".txt"}) == "你好"
    assert extract_text("a.md", "# 标题".encode(), {".md"}) == "# 标题"
    assert extract_text("a.zip", b"x", {".txt"}) == ""


def test_extract_docx():
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("标题一", level=1)
    doc.add_paragraph("正文内容")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.save(buf)
    text = extract_text("a.docx", buf.getvalue(), {".docx"})
    assert "# 标题一" in text
    assert "正文内容" in text
    assert "A | B" in text


def _minimal_pdf(text: str) -> bytes:
    stream = ("BT /F1 12 Tf 72 720 Td (%s) Tj ET" % text).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += ("%d 0 obj\n" % i).encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += ("xref\n0 %d\n" % (len(objects) + 1)).encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)).encode()
    return bytes(out)


def test_extract_pdf():
    text = extract_text("a.pdf", _minimal_pdf("Hello PDF"), {".pdf"})
    assert "Hello PDF" in text


# ── clean ───────────────────────────────────────────────────────────
def test_clean_universal():
    text = "A\u3000B\r\nC\x00D\n\n\n\nE"
    assert clean_text(text, kind="markdown") == "A B\nCD\n\nE"


def test_clean_pdf_noise_and_hyphen():
    text = "正文\nPage 3 of 10\n第5页\nhyphen-\nated word\n"
    cleaned = clean_text(text, kind="pdf", clean_pdf=True)
    assert "Page 3" not in cleaned
    assert "第5页" not in cleaned
    assert "hyphenated" in cleaned


def test_clean_join_broken_lines():
    text = "这是第一行\n这是续行\n\n第二段\n"
    cleaned = clean_text(text, kind="pdf", clean_pdf=True)
    assert "这是第一行 这是续行" in cleaned
    assert "第二段" in cleaned


def test_clean_code_preserves_indent():
    code = "def f():\n    return 1\n"
    assert "    return 1" in clean_text(code, kind="code")


# ── chunker ─────────────────────────────────────────────────────────
def test_kind_for_ext():
    assert kind_for_ext(".md") == "markdown"
    assert kind_for_ext(".docx") == "markdown"
    assert kind_for_ext(".py") == "code"
    assert kind_for_ext(".txt") == "text"


def test_chunk_markdown_heading_path():
    md = "# 安装\n\n## 依赖\n\n需要安装 numpy。\n\n## 用法\n\n运行命令。\n"
    chunks = chunk_text(md, ".md", chunk_size=50, overlap=5, hard_max=200)
    headings = {c.heading for c in chunks}
    assert "# 安装 > ## 依赖" in headings
    assert "# 安装 > ## 用法" in headings
    # 块文本包含标题前缀
    for c in chunks:
        if c.heading:
            assert c.text.startswith(c.heading)


def test_chunk_code_split_defs():
    code = "def a():\n    return 1\n\n\ndef b():\n    return 2\n"
    chunks = chunk_text(code, ".py", chunk_size=6, overlap=0, hard_max=100)
    assert len(chunks) == 2
    assert "def a" in chunks[0].text and "def b" not in chunks[0].text
    assert "def b" in chunks[1].text


def test_chunk_chinese_not_split():
    # 无标点长句不拆，整段作为一个单元保留
    text = "这是一段没有任何标点的超长中文内容" * 50
    chunks = chunk_text(text, ".txt", chunk_size=20, overlap=0, hard_max=10000)
    joined = "".join(c.text for c in chunks)
    assert joined.count("这是一段没有任何标点的超长中文内容") == 50


def test_chunk_plain_multiple():
    text = "。".join("第%d句内容" % i for i in range(100))
    chunks = chunk_text(text, ".txt", chunk_size=20, overlap=5, hard_max=200)
    assert len(chunks) > 1


def test_dedup_chunks():
    chunks = [Chunk("a"), Chunk("a"), Chunk("b"), Chunk("   ")]
    out = dedup_chunks(chunks)
    assert [c.text for c in out] == ["a", "b"]


# ── store / service ─────────────────────────────────────────────────
def test_store_dedup_and_delete(tmp_path):
    store = RagStore(str(tmp_path / "index.sqlite3"), str(tmp_path / "docs"))
    doc_id = store.add_document("a.txt", "sha", 3, ".txt", source_path="/tmp/x")
    assert store.get_by_sha256("sha") == doc_id
    store.add_chunks(doc_id, [Chunk(text="hello", token_count=1)])
    assert store.count_chunks(doc_id) == 1
    assert store.delete_document(doc_id) == "/tmp/x"
    assert store.get_by_sha256("sha") is None
    store.close()


def test_service_upload_dedup(tmp_path):
    cfg = RagConfig(docs_dir=str(tmp_path / "docs"), index_file=str(tmp_path / "index.sqlite3"))
    svc = RagService(cfg)
    r1 = asyncio.run(svc.upload("a.txt", "这是一段测试内容。".encode()))
    assert r1["skipped"] is False and r1["chunks"] >= 1 and r1["doc_id"]
    r2 = asyncio.run(svc.upload("a.txt", "这是一段测试内容。".encode()))
    assert r2["skipped"] is True and r2["doc_id"] == r1["doc_id"]
    assert len(svc.list_documents()) == 1
    assert svc.delete(r1["doc_id"]) is True
    assert svc.list_documents() == []


# ── Phase 2：BM25 / rerank / 检索 ───────────────────────────────────
def test_bm25_tokenize_and_search():
    assert tokenize("Hello 世界 world") == ["hello", "世", "界", "world"]
    idx = BM25Index()
    idx.build(["苹果是一种水果", "篮球是一项运动", "苹果手机很好用"])
    hits = [i for i, _ in idx.search("苹果")]
    assert 0 in hits and 2 in hits


def test_reranker_parse_order():
    assert _parse_order("[2, 0, 1]", 3) == [2, 0, 1]
    assert _parse_order("garbage", 3) == [0, 1, 2]
    assert _parse_order("[0]", 3) == [0, 1, 2]


class _FakeEmbedder:
    available = True

    async def embed(self, texts):
        return [[1.0, 0.0, 0.0, 0.0] for _ in texts]


class _FakeReranker:
    async def rerank(self, query, chunks):
        return list(range(len(chunks)))


def test_service_search_bm25(tmp_path):
    cfg = RagConfig(docs_dir=str(tmp_path / "docs"), index_file=str(tmp_path / "index.sqlite3"))
    svc = RagService(cfg, embedder=_FakeEmbedder(), reranker=_FakeReranker())
    asyncio.run(svc.upload("a.txt", "苹果是一种水果。\n\n天空是蓝色的。".encode()))
    results = asyncio.run(svc.search("苹果"))
    assert any("苹果" in r["text"] for r in results)


def test_service_search_empty(tmp_path):
    cfg = RagConfig(docs_dir=str(tmp_path / "docs"), index_file=str(tmp_path / "index.sqlite3"))
    svc = RagService(cfg)
    assert asyncio.run(svc.search("什么都不存在的查询")) == []


class _BoomEmbedder:
    available = True

    async def embed(self, texts):
        raise RuntimeError("embeddings endpoint unavailable")


def test_service_upload_embed_failure_falls_back(tmp_path):
    cfg = RagConfig(docs_dir=str(tmp_path / "docs"), index_file=str(tmp_path / "index.sqlite3"))
    svc = RagService(cfg, embedder=_BoomEmbedder(), reranker=_FakeReranker())
    r = asyncio.run(svc.upload("a.txt", "苹果是一种水果。".encode()))
    assert r["skipped"] is False and r["chunks"] >= 1
    results = asyncio.run(svc.search("苹果"))
    assert any("苹果" in x["text"] for x in results)


def test_service_upload_corrupt_file_skips(tmp_path):
    cfg = RagConfig(docs_dir=str(tmp_path / "docs"), index_file=str(tmp_path / "index.sqlite3"))
    svc = RagService(cfg)
    r = asyncio.run(svc.upload("bad.pdf", b"not a pdf at all"))
    assert r["skipped"] is True and "error" in r
