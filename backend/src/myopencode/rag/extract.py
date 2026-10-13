"""文本提取：按扩展名把文件字节抽成纯文本。

- PDF：pypdf（内存解析，不落盘）
- Word(.docx)：python-docx，标题样式映射为 Markdown 标题，附带表格文本
- 其它明文/代码：utf-8 直读（errors=replace）
"""
from __future__ import annotations

import io
from typing import Set


def file_ext(name: str) -> str:
    """取小写扩展名（含点）；无扩展名返回空串。"""
    dot = name.rfind(".")
    return name[dot:].lower() if dot != -1 else ""


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = _heading_level(style)
            lines.append("#" * level + " " + text)
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n\n".join(lines).strip()


def _heading_level(style: str) -> int:
    suffix = style.replace("heading", "").strip()
    try:
        level = int(suffix)
    except ValueError:
        level = 1
    return max(1, min(level, 6))


def extract_text(name: str, data: bytes, allowed_extensions: Set[str]) -> str:
    """按扩展名抽取文本；不在允许列表内返回空串。"""
    ext = file_ext(name)
    if ext not in allowed_extensions:
        return ""
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    return data.decode("utf-8", errors="replace")